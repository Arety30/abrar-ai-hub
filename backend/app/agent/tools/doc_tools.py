"""Offline document generation: PDF, PowerPoint, spreadsheets + local vision.

These replace the old online Google Docs/Sheets/Slides tools with fully local
equivalents (fpdf2 / python-pptx / openpyxl) — no internet, no OAuth.
"""
from __future__ import annotations

import json

from ..registry import registry
from ..safety import workspace_out, resolve_path
from ...providers import ollama


@registry.tool(
    "create_pdf",
    "Generate a PDF document offline from a title and body text.",
    {
        "filename": {"type": "string", "description": "Output file name, e.g. report.pdf"},
        "title": {"type": "string", "description": "Document title"},
        "body": {"type": "string", "description": "Body text (newlines become paragraphs)"},
    },
    danger=True, plugin="pdf",
)
def create_pdf(filename: str, title: str = "", body: str = "") -> str:
    from fpdf import FPDF

    if not filename.lower().endswith(".pdf"):
        filename += ".pdf"
    out = workspace_out(filename)
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.multi_cell(0, 10, title or "Document")
    pdf.ln(4)
    pdf.set_font("Helvetica", "", 12)
    for para in (body or "").split("\n"):
        pdf.multi_cell(0, 7, para)
        pdf.ln(1)
    pdf.output(str(out))
    return f"Created PDF: {out}"


@registry.tool(
    "read_pdf",
    "Extract text from a local PDF file.",
    {"path": {"type": "string", "description": "Path to the PDF"}},
    plugin="pdf",
)
def read_pdf(path: str) -> str:
    from pypdf import PdfReader

    p = resolve_path(path, must_exist=True)
    reader = PdfReader(str(p))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return text[:20000] or "(no extractable text)"


@registry.tool(
    "create_slides",
    "Generate a PowerPoint (.pptx) deck offline. Pass slides as a JSON list of "
    '{"title": str, "bullets": [str, ...]}.',
    {
        "filename": {"type": "string", "description": "Output file, e.g. deck.pptx"},
        "slides_json": {"type": "string", "description": 'JSON list of {"title","bullets"}'},
    },
    danger=True, plugin="slides",
)
def create_slides(filename: str, slides_json: str = "[]") -> str:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    if not filename.lower().endswith(".pptx"):
        filename += ".pptx"
    out = workspace_out(filename)
    try:
        slides = json.loads(slides_json)
        assert isinstance(slides, list)
    except Exception:
        return "slides_json must be a JSON list of {title, bullets}."

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for s in slides:
        layout = prs.slide_layouts[1]  # title + content
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = str(s.get("title", ""))
        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = s.get("bullets") or []
        for i, b in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(b)
            para.font.size = Pt(20)
    prs.save(str(out))
    return f"Created deck with {len(slides)} slides: {out}"


@registry.tool(
    "create_spreadsheet",
    "Generate an Excel (.xlsx) spreadsheet offline. Pass rows as a JSON list of lists.",
    {
        "filename": {"type": "string", "description": "Output file, e.g. data.xlsx"},
        "rows_json": {"type": "string", "description": "JSON list of rows (each a list of cells)"},
        "sheet_name": {"type": "string", "description": "Sheet name", "_optional": True},
    },
    danger=True, plugin="sheets",
)
def create_spreadsheet(filename: str, rows_json: str = "[]", sheet_name: str = "Sheet1") -> str:
    from openpyxl import Workbook

    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    out = workspace_out(filename)
    try:
        rows = json.loads(rows_json)
        assert isinstance(rows, list)
    except Exception:
        return "rows_json must be a JSON list of lists."

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Sheet1")[:31]
    for row in rows:
        ws.append(row if isinstance(row, list) else [row])
    wb.save(str(out))
    return f"Created spreadsheet ({len(rows)} rows): {out}"


@registry.tool(
    "generate_presentation",
    "Generate a styled PowerPoint deck with a title slide and speaker notes. Pass "
    'slides as a JSON list of {"title","bullets":[...],"notes":"speaker notes"}.',
    {
        "filename": {"type": "string", "description": "Output file, e.g. talk.pptx"},
        "deck_title": {"type": "string", "description": "Title-slide heading"},
        "slides_json": {"type": "string", "description": 'JSON list of {title,bullets,notes}'},
    },
    danger=True, plugin="slides",
)
def generate_presentation(filename: str, deck_title: str, slides_json: str = "[]") -> str:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    if not filename.lower().endswith(".pptx"):
        filename += ".pptx"
    out = workspace_out(filename)
    try:
        slides = json.loads(slides_json)
        assert isinstance(slides, list)
    except Exception:
        return "slides_json must be a JSON list of {title,bullets,notes}."

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    accent = RGBColor(0xEF, 0x44, 0x44)

    # title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = deck_title
    if title_slide.placeholders[1]:
        title_slide.placeholders[1].text = "Generated offline by Abrar AI Hub"

    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(s.get("title", ""))
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = accent
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, b in enumerate(s.get("bullets") or []):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(b)
            para.font.size = Pt(20)
        notes = s.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)
    prs.save(str(out))
    return f"Created presentation '{deck_title}' with {len(slides)} slides (+ speaker notes): {out}"


@registry.tool(
    "analyze_image",
    "Describe or answer a question about a local image using the offline vision model.",
    {
        "path": {"type": "string", "description": "Path to an image file"},
        "question": {"type": "string", "description": "What to ask about the image", "_optional": True},
    },
    plugin="vision",
)
def analyze_image(path: str, question: str = "Describe this image in detail.") -> str:
    p = resolve_path(path, must_exist=True)
    return ollama.vision(question or "Describe this image in detail.", [str(p)])
